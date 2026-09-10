from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/AutoConnect Testing Guide.docx")


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def style_run(run, size=10, bold=False, color="1F2937"):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(cell, value, bold=False, color="1F2937", align=None):
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.space_before = Pt(2)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(value)
    style_run(run, 9, bold, color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_border(cell)


def add_table(doc, rows, widths):
    table = doc.add_table(rows=1, cols=len(widths))
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0].cells
    for index, value in enumerate(rows[0]):
        header[index].width = Inches(widths[index])
        set_cell_fill(header[index], "0F766E")
        add_text(header[index], value, True, "FFFFFF")
    for row_index, values in enumerate(rows[1:]):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].width = Inches(widths[index])
            if row_index % 2:
                set_cell_fill(cells[index], "F0FDFA")
            add_text(cells[index], value, align=WD_ALIGN_PARAGRAPH.CENTER if index == 0 else None)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    style_run(r, 16 if level == 1 else 12, True, "000000")
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead:
        r = p.add_run(bold_lead)
        style_run(r, 10, True, "000000")
    r = p.add_run(text)
    style_run(r, 10, False, "374151")
    return p


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(7)
    r = title.add_run("AutoConnect Site Testing Guide")
    style_run(r, 25, True, "000000")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    r = subtitle.add_run("Simple checks for buyers sellers providers and administrators")
    style_run(r, 12, False, "374151")

    add_body(doc, "Use this guide before launch. Start with the public pages without signing in. Later, when test logins are created, repeat the account checks in the role sections. Mark each check as Pass, Fail, or Needs partner setup.", "Main point: ")
    add_body(doc, "Do not use a real bank transfer, card, M Pesa payment, or personal document while testing. Payment remains pending until a real provider confirms it.", "Important: ")

    add_heading(doc, "How to test", 1)
    rows = [
        ["Step", "What to do", "Expected result"],
        ["1", "Open the site on a phone and a computer.", "The page fits the screen. Text and buttons are readable. No horizontal scrolling on normal pages."],
        ["2", "Click the links in the top menu and footer.", "Each link opens the right page or a clear error page. No blank page or dead button."],
        ["3", "Try each main action once.", "The action works, asks you to sign in, or explains why it cannot continue. It must not pretend it succeeded."],
        ["4", "If something fails, take a screenshot and copy the page link.", "Report the exact step and message. Do not report only that it is broken."],
    ]
    add_table(doc, rows, [0.45, 2.45, 3.55])

    add_heading(doc, "Public buyer journey", 1)
    rows = [
        ["Test", "What to do", "Expected result"],
        ["Home", "Open Home. Use search, Browse cars, Parts, Services, Import and Auctions.", "Each action opens the correct area. Cards have images where real images exist and empty states explain what is missing."],
        ["Search", "Search by make or model. Change filters and clear them.", "The vehicle list changes or clearly says there are no matching cars. It never hides a loading or database error as no results."],
        ["Vehicle", "Open a vehicle. Check photos, facts, evidence, seller contact, compare and viewing request.", "Facts match the listing. Missing evidence is marked missing. Enquiry and viewing forms ask for required information."],
        ["Compare", "Add two cars to Compare. Remove one. Try a single car.", "Comparison shows side by side facts. Empty and single car states explain what to do next."],
        ["Dealership or yard", "Open a dealership or yard and browse its stock.", "The stock belongs to that seller or yard. It must not silently show an unrelated general catalogue."],
        ["Parts", "Open a part, choose a garage vehicle if signed in, and request a quote.", "The request is sent only after sign in and required details. It says quote requested, not purchased or delivered."],
        ["Services", "Open Services and request an appointment.", "The form asks for a vehicle, provider, date and notes. After sending, it says the provider can quote or confirm."],
        ["Import", "Use the import estimate and submit a sourcing request.", "Costs are clearly estimates. A request is recorded; it is not a confirmed vehicle order."],
        ["Import tracker", "Search for an existing order reference.", "It shows only data attached to that order. Missing carrier or clearing updates are shown as missing."],
        ["Auctions", "Open Auctions. Choose an offer and inspect the end time, price and bids.", "Only admin approved offers are public. Bidder names are masked. A bid or winning reservation requires sign in."],
    ]
    add_table(doc, rows, [0.9, 2.55, 3.0])

    add_heading(doc, "Signed in buyer journey", 1)
    rows = [
        ["Test", "What to do", "Expected result"],
        ["Account", "Open Account and change only test profile details.", "The changes remain after refresh. Other users cannot see private information."],
        ["My Garage", "Add a test vehicle. Add mileage, a reminder, a service booking and a part request.", "Every item stays attached to the selected vehicle. A signed in user never sees the preview owner data as their own."],
        ["Booking", "Ask for service. As provider, quote it. As buyer, approve it. Mark work complete and leave a review.", "Status moves request to quote to approval to completed. A receipt appears after completion, not before."],
        ["Auction", "Place a valid bid on a live test auction. Let it end. Claim the win.", "The bid shows under a masked alias. Claiming makes a time limited reservation with payment pending, never paid."],
        ["Purchase", "Open My purchases and a transaction.", "The status matches the real transaction record. A bank reference or reservation alone must not show payment confirmed."],
        ["Dark mode", "Turn on dark mode and repeat a few pages.", "Text, form labels, buttons and warnings remain readable with clear contrast."],
    ]
    add_table(doc, rows, [0.9, 2.55, 3.0])

    add_heading(doc, "Seller journey", 1)
    rows = [
        ["Test", "What to do", "Expected result"],
        ["Seller access", "Open Seller workspace with a test seller.", "An unapproved seller sees what is blocked. An approved seller sees their own tools only."],
        ["Listing", "Create a test listing with photos and documents.", "The listing is saved as pending review. It is not public until approved by an administrator."],
        ["Evidence", "Submit a logbook, inspection or history evidence link.", "It stays private and pending until an administrator reviews it."],
        ["Enquiries", "Reply to a buyer test enquiry and arrange a viewing.", "The message stays in the correct enquiry thread and is visible to the right buyer and seller."],
        ["Auction", "Create a timed auction or flash offer for an approved test vehicle.", "It enters pending approval. It cannot appear publicly until an administrator approves it."],
        ["Reservation", "Review an auction reservation after the auction ends.", "The status is payment pending. The seller cannot treat it as cleared payment."],
    ]
    add_table(doc, rows, [0.9, 2.55, 3.0])

    add_heading(doc, "Provider and administrator journey", 1)
    rows = [
        ["Test", "What to do", "Expected result"],
        ["Provider booking", "Open service bookings as the assigned provider. Send a quote and complete a booking.", "Only the assigned provider can manage it. The buyer sees each status change."],
        ["Listing review", "Approve and reject separate test listings.", "The public list changes only after approval. A rejection contains a usable reason."],
        ["Evidence review", "Open Vehicle evidence. Verify one test item and reject another.", "Verified public evidence can be shown on its vehicle. Rejected evidence stays unavailable to buyers."],
        ["Auction review", "Open Auction review. Approve one test auction, reject one, then cancel one test auction.", "Only approved auctions become public. The audit trail records the decision. Do not cancel a real auction while testing."],
        ["Expired reservations", "Run Process expired reservations using expired test data only.", "Only due payment pending reservations change to expired. Buyer and seller notifications are created."],
        ["Disputes", "Create a test dispute from a transaction and review it as admin.", "The dispute has a clear status, notes and next owner. It does not silently close."],
    ]
    add_table(doc, rows, [0.9, 2.55, 3.0])

    add_heading(doc, "Failure checks", 1)
    rows = [
        ["Test", "What to do", "Expected result"],
        ["No sign in", "Try bid, booking, parts quote, evidence submission and payment action while signed out.", "The site asks for sign in. It does not create a record under another user."],
        ["Bad form", "Leave required fields empty or enter an invalid value.", "The page explains what must be fixed and keeps the entered information where possible."],
        ["Slow internet", "Use normal phone data or browser throttling. Refresh during a load.", "Loading state is visible. A failed request shows an error and retry option, not an empty successful screen."],
        ["No data", "Use a search that has no match.", "The page says no matching records and provides a next action."],
        ["Permission", "Try to open another user's private account link if you have one.", "Access is denied or the record is not found. Private records are never displayed."],
    ]
    add_table(doc, rows, [0.9, 2.55, 3.0])

    add_heading(doc, "What is ready and what still needs a partner", 1)
    add_body(doc, "Ready in the code: marketplace browsing, enquiries, service booking stages, parts quote requests, seller submission, evidence review, auction approvals, bid privacy, reservation expiry, audit logs, and signed in funnel events.")
    add_body(doc, "Needs a real external setup: verified payment confirmations, a named bank or licensed custody arrangement, PesaPal or Stripe account configuration, real carrier and clearing updates, provider account onboarding, external error monitoring, uptime alerts, automated backup verification and a formal mobile end to end test run.")
    add_body(doc, "The site should not say a payment is held, escrowed or complete until the actual payment provider or bank evidence has been checked.", "Rule for testers: ")

    add_heading(doc, "Test report template", 1)
    rows = [
        ["Field", "Write this"],
        ["Date and tester", "Example: 10 September 2026, Jane"],
        ["Device and browser", "Example: Samsung A54, Chrome, mobile data"],
        ["Page link", "Copy the exact page address"],
        ["Steps", "Example: Open Cars, choose Toyota, open first result, press Request viewing"],
        ["Expected result", "Copy the expected result from this guide"],
        ["Actual result", "What actually happened, including any message"],
        ["Evidence", "Screenshot or short screen recording"],
        ["Priority", "Blocker, high, normal or low"],
    ]
    add_table(doc, rows, [1.6, 4.85])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("AutoConnect Testing Guide")
    style_run(fr, 8, False, "6B7280")
    doc.save(OUT)


if __name__ == "__main__":
    build()
